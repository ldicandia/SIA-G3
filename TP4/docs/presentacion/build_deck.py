"""Presentación técnica universitaria, guiada por resultados y notas del expositor.

ENERGY 1 / RHYTHM 2 / MOTION 1: lectura calma, composición según la evidencia,
sin animaciones. Blanco y Arial favorecen la proyección; el acento verde petróleo
señala la condición elegida. Figuras grandes y espacio libre separan observación,
comparación y decisión. La estructura sigue el orden del análisis, sin tarjetas.
"""

from pathlib import Path
import json
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.patches import RegularPolygon
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT = HERE / "TP4-presentacion.pptx"
INK = "#18242C"
TEAL = "#176B72"
COPPER = "#A25D35"
MUTED = "#5E6B73"
LIGHT = "#DDE3E5"
FONT = "Arial"
TOPICS = {"football": "Fútbol", "makeup": "Belleza", "dogs": "Perros", "art": "Arte"}
QR_NAMES = {"food": "Principal", "wine": "Vinos", "beverages": "Bebidas", "dessert": "Postres", "specials": "Especiales"}
plt.rcParams.update({"font.family": FONT, "font.size": 15, "axes.labelsize": 16,
                     "axes.titlesize": 17, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.edgecolor": MUTED, "axes.labelcolor": INK, "text.color": INK,
                     "xtick.color": INK, "ytick.color": INK, "figure.facecolor": "white",
                     "axes.facecolor": "white", "savefig.facecolor": "white"})


def summary(name):
    return json.loads((ROOT / "runs" / name / "summary.json").read_text())


def data(name, filename):
    return pd.read_csv(ROOT / "runs" / name / filename)


def finish(fig, name):
    fig.tight_layout(pad=1.2)
    path = FIG / f"{name}.png"
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def axis(ax):
    ax.grid(axis="y", color=LIGHT, linewidth=.7, zorder=0)
    ax.set_axisbelow(True)


def point_plot(ax, rows, metric, labels, ylabel, color=TEAL):
    x = np.arange(len(rows))
    ax.errorbar(x, rows[f"{metric}_mean"], yerr=rows[f"{metric}_std"], fmt="o-",
                color=color, capsize=5, linewidth=2, markersize=8)
    ax.set_xticks(x, labels=labels)
    ax.set_ylabel(ylabel)
    ax.margins(x=.16)
    axis(ax)


def charts():
    FIG.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / "runs/hopfield_letters/storage.png", FIG / "hopfield_letters__storage.png")
    raw = pd.read_csv(ROOT / "datos/creadores.csv")
    top = raw.nlargest(3, "followers").iloc[::-1]
    fig, ax = plt.subplots(figsize=(11.6, 4.4))
    left = np.zeros(3)
    colors = [TEAL, COPPER, "#8A999F", "#B8C3C8"]
    for (topic, label), color in zip(TOPICS.items(), colors):
        values = top[f"topic_{topic}_pct"].to_numpy()
        ax.barh(top["name"], values, left=left, color=color, label=label, height=.55)
        for y, (start, width) in enumerate(zip(left, values)):
            if width >= 15:
                ax.text(start + width / 2, y, f"{width:.0f}%", ha="center", va="center",
                        color="white" if color in [TEAL, COPPER] else INK, fontsize=17)
        left += values
    ax.set_xlim(0, 100)
    ax.set_xlabel("Contenido de cada tópico (%)")
    ax.legend(ncol=4, frameon=False, bbox_to_anchor=(.5, 1.18), loc="upper center")
    ax.spines[["left", "bottom"]].set_visible(False)
    finish(fig, "followers_topics")

    fig, axes = plt.subplots(1, 3, figsize=(12.8, 4.3))
    for ax, col, label in zip(axes, ["followers", "price_usd", "comments_per_1000"],
                              ["Seguidores", "Precio por publicación (USD)", "Comentarios por 1000"]):
        ax.hist(raw[col], bins=10, color=TEAL, edgecolor="white")
        ax.set_xlabel(label)
        axis(ax)
    axes[0].set_ylabel("Creadores")
    finish(fig, "distributions")

    variance = data("pca", "explained_variance.csv")
    ratios = variance["explained_variance_ratio"].to_numpy() * 100
    fig, ax = plt.subplots(figsize=(11.6, 4.8))
    x = np.arange(1, len(ratios) + 1)
    ax.bar(x, ratios, color=LIGHT, label="Por componente")
    ax.plot(x, np.cumsum(ratios), "o-", color=TEAL, markersize=4, linewidth=2, label="Acumulada")
    for k in [2, 6, 10]:
        value = np.cumsum(ratios)[k - 1]
        label = f"{value:.1f}%".replace(".", ",")
        ax.scatter([k], [value], s=65, color=TEAL)
        ax.annotate(label, (k, value), (k + .8, value - 8), fontsize=18, color=TEAL)
    ax.set(xlim=(.3, 27.7), ylim=(0, 105), xlabel="Componentes conservadas", ylabel="Varianza explicada (%)")
    ax.set_xticks([1, 2, 6, 10, 15, 20, 27])
    ax.legend(frameon=False, loc="lower right")
    axis(ax)
    finish(fig, "pca_variance")

    scores = data("pca", "scores.csv")
    loadings = data("pca", "loadings.csv").set_index("feature")
    fig, ax = plt.subplots(figsize=(11.4, 5.8))
    ax.scatter(scores.PC1, scores.PC2, color="#87959C", s=45, alpha=.8, zorder=2)
    arrows = {"average_caption_words": "Epígrafe", "average_impressions": "Impresiones",
              "content_showcase_pct": "Showcase", "topic_art_pct": "Arte",
              "argentina_audience_pct": "Audiencia argentina"}
    for feature, label in arrows.items():
        xy = loadings.loc[feature, ["PC1", "PC2"]].to_numpy(dtype=float) * 9
        ax.annotate("", xy, (0, 0), arrowprops={"arrowstyle": "->", "color": TEAL, "lw": 1.6})
        ax.text(xy[0] * 1.12, xy[1] * 1.12, label, fontsize=14, color=TEAL,
                ha="left" if xy[0] >= 0 else "right", va="center")
    for creator in ["@mateconideas", "@mundoanimal.ar"]:
        r = scores.loc[scores.name == creator].iloc[0]
        ax.scatter([r.PC1], [r.PC2], s=70, color=COPPER)
        ax.annotate(creator, (r.PC1, r.PC2), (7, 7), textcoords="offset points", fontsize=13)
    ax.axhline(0, color=LIGHT, lw=1); ax.axvline(0, color=LIGHT, lw=1)
    ax.set(xlabel="PC1 (49,7%)", ylabel="PC2 (9,2%)", xlim=(-8.5, 8.5), ylim=(-4.6, 5.4))
    finish(fig, "pca_biplot")

    neurons = data("som", "neurons.csv")
    fig, ax = plt.subplots(figsize=(7, 5))
    for r in neurons.itertuples():
        distance = np.hypot(r.x - neurons.loc[5, "x"], r.y - neurons.loc[5, "y"])
        color = TEAL if r.neuron == 5 else "#C5DCDC" if distance <= 1.00001 else "white"
        ax.add_patch(RegularPolygon((r.x, r.y), 6, radius=1 / np.sqrt(3), facecolor=color, edgecolor=MUTED))
        ax.text(r.x, r.y, str(r.neuron), ha="center", va="center", fontsize=18,
                color="white" if r.neuron == 5 else INK)
    ax.set(xlim=(-.8, 4.2), ylim=(3.3, -.8), aspect="equal"); ax.axis("off")
    finish(fig, "som_neighborhood")

    table = data("som_sweep", "comparison.csv")
    for parameter, name, metrics, ylabel in [
        ("params.eta0", "som_eta", ["mean_bmu_distance", "u_mean"], ["Distancia al ganador", "U media"]),
        ("params.radius0", "som_radius", ["mean_bmu_distance", "empty_neuron_frac"], ["Distancia al ganador", "Fracción vacía"]),
        ("params.rows", "som_size", ["mean_bmu_distance", "n_empty_neurons"], ["Distancia al ganador", "Neuronas vacías"]),
    ]:
        rows = table.loc[table.varied_parameter == parameter]
        labels = [f"{x:g}" for x in rows.parameter_value]
        if parameter == "params.rows":
            labels = [f"{x}×{x}" for x in labels]
        fig, axes = plt.subplots(1, 2, figsize=(11.4, 4.9))
        for ax, metric, label in zip(axes, metrics, ylabel):
            point_plot(ax, rows, metric, labels, label)
            ax.set_xlabel({"params.eta0": r"Tasa inicial $\eta_0$", "params.radius0": "Radio constante", "params.rows": "Grilla"}[parameter])
        finish(fig, name)

    design = data("som_design_sweep", "comparison.csv")
    fig, axes = plt.subplots(1, 3, figsize=(12.8, 4.4))
    for ax, frame, parameter, title in [(axes[0], table, "params.epochs", "Épocas"),
                                      (axes[1], design, "params.topology", "Topología"),
                                      (axes[2], design, "params.init", "Inicialización")]:
        rows = frame.loc[frame.varied_parameter == parameter]
        labels = [str(value) for value in rows.parameter_value]
        labels = [{"50.0": "50", "200.0": "200", "sample": "Muestras", "random": "Uniforme"}.get(value, value) for value in labels]
        point_plot(ax, rows, "mean_bmu_distance", labels, "Distancia al ganador" if ax == axes[0] else "")
        ax.set_title(title)
    finish(fig, "som_design")

    for field, name in [("hits", "som_hits"), ("u_distance", "som_u")]:
        fig, ax = plt.subplots(figsize=(8.4, 5.5))
        cmap = LinearSegmentedColormap.from_list("map", ["#EFF4F4", TEAL])
        lo, hi = neurons[field].min(), neurons[field].max()
        for r in neurons.itertuples():
            value = getattr(r, field)
            normalized = (value - lo) / max(hi - lo, 1e-9)
            ax.add_patch(RegularPolygon((r.x, r.y), 6, radius=1 / np.sqrt(3),
                                        facecolor=cmap(normalized), edgecolor="white", linewidth=2))
            label = str(int(value)) if field == "hits" else f"{value:.2f}"
            ax.text(r.x, r.y, label, ha="center", va="center", fontsize=19,
                    color="white" if normalized > .55 else INK, fontweight="bold")
            ax.text(r.x, r.y + .33, f"{r.neuron}", ha="center", va="center", fontsize=10,
                    color="white" if normalized > .55 else INK)
        if field == "u_distance":
            edge = data("som", "u_edges.csv").iloc[0]
            a, b = neurons.iloc[int(edge.neuron_a)], neurons.iloc[int(edge.neuron_b)]
            ax.plot([a.x, b.x], [a.y, b.y], color=COPPER, linewidth=4, zorder=4)
        ax.set(xlim=(-.8, 4.2), ylim=(3.4, -.8), aspect="equal"); ax.axis("off")
        finish(fig, name)

    weight = data("oja", "weights_stats.csv").set_index("feature")
    show = ["average_caption_words", "average_impressions", "average_reach", "saves_per_1000", "content_tutorial_pct", "price_usd"]
    labels = ["Epígrafe", "Impresiones", "Alcance", "Guardados", "Tutoriales", "Precio"]
    fig, ax = plt.subplots(figsize=(10.4, 5.4))
    y = np.arange(len(show))
    ax.errorbar(weight.loc[show, "oja_weight_mean"], y - .12,
                xerr=weight.loc[show, "oja_weight_std"], fmt="o", color=TEAL, capsize=4, markersize=9, label="Oja, media ± DE")
    ax.scatter(weight.loc[show, "pc1_loading"], y + .12, marker="x", s=85, linewidths=2, color=COPPER, label="PC1")
    ax.set_yticks(y, labels=labels); ax.invert_yaxis()
    ax.axvline(0, color=LIGHT); ax.set_xlabel("Coeficiente del vector normalizado")
    ax.legend(frameon=False, loc="lower right")
    finish(fig, "oja_weights")

    index = data("oja", "index_stats.csv")
    fig, ax = plt.subplots(figsize=(10, 5.5))
    ax.errorbar(index.pc1_score, index.oja_index_mean, yerr=index.oja_index_std, fmt="o", ms=5,
                color="#8F9CA2", alpha=.8, linewidth=.8, capsize=2)
    stable = index.top5_count == index.n_repeats
    ax.scatter(index.loc[stable, "pc1_score"], index.loc[stable, "oja_index_mean"], s=55, color=TEAL, zorder=3)
    ax.plot([-7.5, 7.5], [-7.5, 7.5], color=INK, lw=1, ls=":")
    ax.set(xlabel="PC1", ylabel="Índice de Oja, media ± DE")
    axis(ax)
    finish(fig, "oja_scores")

    profiles = data("campaigns", "strategy_profiles.csv").query("client == 'football'")
    compare = profiles.loc[profiles.som_bmu.isin([5, 13])].sort_values("som_bmu")
    fig, axes = plt.subplots(1, 3, figsize=(12.2, 4.4))
    for ax, column, label in zip(axes, ["mean_content_tutorial_pct", "mean_saves_per_1000", "mean_followers"],
                                ["Tutoriales (%)", "Guardados por 1000", "Seguidores (miles)"]):
        values = compare[column].to_numpy()
        if column == "mean_followers": values = values / 1000
        ax.bar(["Celda 5", "Celda 13"], values, color=[TEAL, COPPER], width=.55)
        for i, value in enumerate(values): ax.text(i, value + max(values) * .04, f"{value:.1f}", ha="center", fontsize=17)
        ax.set_ylim(0, max(values) * 1.25); ax.set_ylabel(label)
        axis(ax)
    finish(fig, "football_profiles")

    dots = data("hopfield_letters", "dot_products.csv").set_index("letter")
    letters = summary("hopfield_letters")
    fig, axes = plt.subplots(1, 2, figsize=(10.8, 4.5))
    for ax, selected, title in zip(axes, [letters["selected_letters"], letters["worst_letters"]], ["Elegidas", "Más correlacionadas"]):
        matrix = dots.loc[selected, selected].to_numpy(dtype=float) / 25
        ax.imshow(np.abs(matrix), cmap="Greys", vmin=0, vmax=1)
        for i in range(4):
            for j in range(4):
                ax.text(j, i, f"{matrix[i,j]:.2f}", ha="center", va="center", fontsize=20,
                        color="white" if abs(matrix[i, j]) > .56 else INK)
        ax.set_xticks(range(4), selected); ax.set_yticks(range(4), selected); ax.set_title(title)
    finish(fig, "letters_dots")

    repeated = data("hopfield_letters", "recall_repetitions.csv")
    fig, ax = plt.subplots(figsize=(11, 5.1))
    for group, update, color, style, label in [("selected", "async", TEAL, "o-", "G,R,T,V asincrónica"),
                                              ("selected", "sync", TEAL, "s--", "G,R,T,V sincrónica"),
                                              ("worst", "async", COPPER, "o-", "H,M,N,W asincrónica")]:
        rows = repeated[(repeated["set"] == group) & (repeated["update"] == update) & (repeated.level != .4)]
        ax.errorbar(np.rint(rows.level * 25), rows.recovery_rate_mean * 100,
                    yerr=rows.recovery_rate_std * 100, fmt=style, color=color, capsize=4, label=label)
    ax.set(xlabel="Bits invertidos de 25", ylabel="Recuperación exacta (%)", ylim=(-3, 105))
    ax.set_xticks([1, 2, 3, 4, 5, 7, 9, 11]); ax.legend(frameon=False, fontsize=13)
    axis(ax)
    finish(fig, "letters_noise")

    sp = summary("hopfield_letters")
    fig, axes = plt.subplots(1, 3, figsize=(12, 4), gridspec_kw={"width_ratios": [1, 1, 1.5]})
    for ax, rows, title in zip(axes[:2], [sp["spurious_query"], sp["spurious_state"]], ["Consulta", "Punto fijo alcanzado"]):
        matrix = np.array([[char == "#" for char in row] for row in rows])
        ax.imshow(matrix, cmap="binary", vmin=0, vmax=1); ax.set_title(title); ax.axis("off")
    trace = data("hopfield_letters", "energy_trace.csv")
    rows = trace.loc[trace.example == f"spurious_{sp['spurious_letter']}"]
    axes[2].plot(rows.step, rows.energy, color=TEAL, lw=2)
    axes[2].set(xlabel="Visitas a neuronas", ylabel="Energía")
    axis(axes[2]); finish(fig, "letters_spurious")

    cap = data("hopfield_letters", "capacity_repetitions.csv")
    fig, ax = plt.subplots(figsize=(10.8, 4.8))
    ax.errorbar(cap.p, cap.frac_sets_all_stable_mean * 100, yerr=cap.frac_sets_all_stable_std * 100,
                fmt="o-", color=TEAL, capsize=4)
    ax.axvline(3.75, color=COPPER, linestyle=":", linewidth=2, label="0,15 × 25 = 3,75")
    ax.set(xlabel="Letras almacenadas por red", ylabel="Subconjuntos enteramente estables (%)", ylim=(-3, 105))
    ax.set_xticks(cap.p); ax.legend(frameon=False); axis(ax)
    finish(fig, "letters_capacity")

    qr = summary("hopfield_qr")
    fig, axes = plt.subplots(1, 5, figsize=(13, 3.1))
    for ax, name in zip(axes, qr["pattern_names"]):
        original = np.asarray(Image.open(ROOT / "datos/hopfield" / f"{name}_menu_qr.png").convert("L"))
        ax.imshow(original, cmap="gray", vmin=0, vmax=255, interpolation="nearest")
        ax.set_title(QR_NAMES[name]); ax.axis("off")
    finish(fig, "qr_originals")

    pairs = data("hopfield_qr", "pairwise_similarity.csv")
    overlap = data("hopfield_qr", "overlaps.csv").set_index("pattern")
    fig, ax = plt.subplots(figsize=(6.3, 5))
    values = overlap.to_numpy()
    ax.imshow(values, cmap=LinearSegmentedColormap.from_list("similarity", ["white", TEAL]), vmin=0, vmax=1)
    for i in range(5):
        for j in range(5):
            ax.text(j, i, f"{values[i,j]:.2f}", ha="center", va="center", fontsize=18,
                    color="white" if values[i,j] > .68 else INK)
    ax.set_xticks(range(5), [QR_NAMES[name] for name in overlap.columns], rotation=25, ha="right")
    ax.set_yticks(range(5), [QR_NAMES[name] for name in overlap.index])
    finish(fig, "qr_similarity")

    fig, axes = plt.subplots(1, 3, figsize=(12.8, 4.2), gridspec_kw={"width_ratios": [1, 1, 1.35]})
    first = np.asarray(Image.open(ROOT / "datos/hopfield/food_menu_qr.png").convert("L"))
    final = np.array([[char == "#" for char in row] for row in qr["spurious_state"]])
    axes[0].imshow(first, cmap="gray", vmin=0, vmax=255, interpolation="nearest")
    axes[1].imshow(final, cmap="binary", vmin=0, vmax=1, interpolation="nearest")
    for ax, title in zip(axes[:2], ["Original sin ruido", "Estado final"]): ax.set_title(title); ax.axis("off")
    repeated = data("hopfield_qr", "recall_repetitions.csv")
    joint = repeated.loc[repeated["set"] == "joint"]
    axes[2].errorbar(joint.level * 100, joint.mean_iterations_mean, yerr=joint.mean_iterations_std,
                     fmt="o-", color=TEAL, capsize=3, markersize=5)
    axes[2].set(xlabel="Módulos invertidos (%)", ylabel="Barridos que cambian el estado")
    axis(axes[2]); finish(fig, "qr_failure")

    storage = data("hopfield_qr", "storage_control.csv")
    fig, axes = plt.subplots(1, 2, figsize=(11.8, 4.5))
    counts = storage.groupby("n_stored").n_fixed_points.mean()
    axes[0].bar(counts.index, counts.values, color=TEAL, width=.6)
    for count, value in counts.items(): axes[0].text(count, value + .12, f"{value:.0f}/{count}", ha="center", fontsize=19)
    axes[0].set(xlabel="QR almacenados por red", ylabel="Patrones estables", ylim=(0, 3))
    axes[0].set_xticks(range(1, 6)); axis(axes[0])
    for condition, color, label in [("single", TEAL, "Uno por red"), ("joint", COPPER, "Cinco juntos")]:
        rows = repeated.loc[repeated["set"] == condition]
        axes[1].errorbar(rows.level * 100, rows.recovery_rate_mean * 100,
                         yerr=rows.recovery_rate_std * 100, fmt="o-", color=color, capsize=3, label=label)
    axes[1].set(xlabel="Módulos invertidos (%)", ylabel="Recuperación exacta (%)", ylim=(-4, 105))
    axes[1].legend(frameon=False, fontsize=13); axis(axes[1]); finish(fig, "qr_controls")

    fig, ax = plt.subplots(figsize=(11, 4.7))
    for metric, color, label in [("query_decode_rate", TEAL, "Lector directo"), ("final_decode_rate", COPPER, "Lector después de Hopfield")]:
        ax.errorbar(joint.level * 100, joint[f"{metric}_mean"] * 100, yerr=joint[f"{metric}_std"] * 100,
                    fmt="o-", color=color, capsize=4, label=label)
    ax.set(xlabel="Módulos invertidos (%)", ylabel="URL objetivo decodificada (%)", ylim=(-3, 105))
    ax.legend(frameon=False); axis(ax); finish(fig, "qr_reader")


def rgb(value):
    return RGBColor.from_string(value.lstrip("#"))


def text(slide, value, x, y, w, h, size=26, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT; paragraph.font.size = Pt(size)
        paragraph.font.bold = bold; paragraph.font.color.rgb = rgb(color)
        paragraph.space_after = Pt(5)
        if align is not None: paragraph.alignment = align
    return shape


def picture(slide, name, x=.8, y=1.55, w=14.4, h=6.7):
    path = FIG / name
    with Image.open(path) as image: iw, ih = image.size
    factor = min(w / iw, h / ih)
    nw, nh = iw * factor, ih * factor
    return slide.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2), Inches(nw), Inches(nh))


SECTION = ""
KICKER = {"1": "Ejercicio 1: Creadores de contenido",
          "1.1": "Ejercicio 1: Creadores de contenido · 1.1 Representación, estandarización y PCA",
          "1.2": "Ejercicio 1: Creadores de contenido · 1.2 Mapa de Kohonen",
          "1.3": "Ejercicio 1: Creadores de contenido · 1.3 Regla de Oja e índice",
          "1.4": "Ejercicio 1: Creadores de contenido · 1.4 Campañas y cambios de datos",
          "2.1": "Ejercicio 2: Memoria asociativa · 2.1 Hopfield: letras 5×5 y códigos QR"}


def new_slide(prs, title, reminders):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb("#FFFFFF")
    if title and SECTION: text(slide, KICKER[SECTION], .72, .12, 14.5, .3, size=15, color=MUTED)
    text(slide, title, .72, .45, 14.5, .9, size=32, bold=True)
    text(slide, "TP4 · Grupo 3", .72, 8.55, 3, .25, size=12, color=MUTED)
    text(slide, str(len(prs.slides)).zfill(2), 14.7, 8.5, .6, .3, size=13, color=MUTED, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = "\n".join("• " + reminder for reminder in reminders)
    return slide


def caption(slide, value):
    text(slide, value, .85, 7.95, 14.2, .45, size=20, color=MUTED)


def build():
    global SECTION
    charts()
    prs = Presentation(); prs.slide_width = Inches(16); prs.slide_height = Inches(9)
    prs.core_properties.title = "TP4 · Aprendizaje no supervisado"
    prs.core_properties.subject = "Creadores de contenido y memoria asociativa"
    prs.core_properties.author = "Grupo 3 · SIA · ITBA"
    pca, som, oja = (summary(name) for name in ["pca", "som", "oja"])
    letters, qr, campaigns = (summary(name) for name in ["hopfield_letters", "hopfield_qr", "campaigns"])
    source_pca = "Fuente: clase PCA, pp. 7-11, 20-32 y 43-46; runs/pca."
    source_som = "Fuente: clase Kohonen, pp. 24-37; runs/som y barridos OAT."
    source_hop = "Fuente: clase Hopfield, pp. 15, 20, 33-38; runs/hopfield_letters."

    s = new_slide(prs, "", [
        "Dos preguntas del enunciado: elegir creadores y recuperar patrones deteriorados.",
        "Caso 1: adecuar creadores al objetivo, la audiencia y el contenido de cada campaña.",
        "Caso 2: estudiar recuperación de patrones incompletos o ruidosos con memoria asociativa.",
        "Recorrido: representar los datos, comparar configuraciones e interpretar resultados y limitaciones.",
        "Fuente: SIATP42026.pdf, ejercicios 1 y 2; datos originales del TP."])
    text(s, "Aprendizaje\nno supervisado", .85, 2.0, 13.8, 2, size=58, bold=True)
    text(s, "Perfiles de creadores y memoria asociativa", .9, 4.65, 13.5, .7, size=29, color=TEAL)
    text(s, "Trabajo Práctico 4 · SIA · ITBA · 2Q 2026", .9, 6.8, 13.5, .5, size=22, color=MUTED)

    SECTION = ""
    s = new_slide(prs, "Índice", ["Ejercicio 1: creadores de contenido. Ejercicio 2: memoria asociativa de Hopfield."])
    for x, head, items in [
        (.9, "Ejercicio 1 · Creadores de contenido", [
            ("1.1", "Representación, estandarización y PCA"),
            ("1.2", "Mapa de Kohonen"),
            ("1.3", "Regla de Oja e índice"),
            ("1.4", "Campañas y cambios de datos")]),
        (8.4, "Ejercicio 2 · Memoria asociativa", [
            ("2.1", "Hopfield: letras 5×5 y códigos QR 37×37")])]:
        text(s, head, x, 2.0, 6.9, .7, size=28, bold=True, color=TEAL)
        for i, (num, label) in enumerate(items):
            text(s, num, x, 3.1 + i * 1.15, 1.1, .7, size=30, bold=True, color=COPPER)
            text(s, label, x + 1.2, 3.1 + i * 1.15, 5.7, 1.0, size=26)

    SECTION = "1"
    s = new_slide(prs, "Los más seguidos publican contenidos distintos", [
        "La agencia elige principalmente por seguidores; necesita adecuación a cada campaña.",
        "Esta figura usa los tres primeros por seguidores del catálogo original.",
        "Los porcentajes corresponden al contenido publicado, no a conversiones ni composición temática de audiencia.",
        "Clientes: refugios, ropa de fútbol, belleza y materiales de arte.",
        "Fuente: enunciado 1, datos/creadores.csv; runs/campaigns/comparison.csv."])
    picture(s, "followers_topics.png", y=2.0, h=5.4)

    SECTION = "1.1"
    s = new_slide(prs, "Elegimos una representación de 27 variables", [
        "50 creadores, 31 variables numéricas; nombres e ID solo identifican filas. Sin faltantes; conservamos los 50 registros.",
        "price_ars duplica price_usd por conversión fija; organic_posts_pct complementa sponsored_posts_pct.",
        "Engagement total suma cuatro interacciones; alcance argentino estimado combina alcance y porcentaje argentino. Conservamos sus componentes.",
        "Retenemos completos los cuatro tópicos y los cuatro formatos para tratarlos simétricamente; las dependencias no impiden PCA.",
        "Las cuatro columnas excluidas siguen en el CSV para interpretar campañas.",
        "Fuente: enunciado pp. 1-4; configs/feature_set.json; runs/eda/identities.csv y feature_set.json."])
    text(s, "31", 1.1, 2.1, 3.8, 1.8, size=112, color=MUTED)
    text(s, "variables numéricas", 1.15, 4.2, 4.7, .7, size=28)
    text(s, "→", 6.05, 2.45, 2, 1.4, size=70, color=MUTED)
    text(s, "27", 9.2, 2.1, 4, 1.8, size=112, color=TEAL)
    text(s, "variables de entrada", 9.25, 4.2, 5.5, .7, size=28)
    text(s, "50 creadores conservados", 1.15, 6.5, 12, .7, size=30)

    s = new_slide(prs, "Estandarizamos antes de comparar perfiles", [
        "EDA: inspeccionamos distribuciones, correlaciones y boxplots en las 31 variables; aquí tres escalas ilustrativas.",
        "Boxplots: 33 celdas en 25 creadores fuera de bigotes de 1,5 RIC. Se señalan, sin quitar ni recortar filas.",
        "Media y desvío se calculan en el catálogo. Z-score con ddof=0, consistente con StandardScaler.",
        "Sin escalar, PC1 explica 96% y predomina seguidores; las unidades dominan la varianza.",
        "Fuente: clase Kohonen p. 19; PCA pp. 25-27; runs/eda; runs/pca_scaling_sweep/standardize-false."])
    picture(s, "distributions.png", y=1.75, h=4.8)
    text(s, "z = (x − μ) / σ", .95, 7.0, 8, .65, size=34, color=TEAL)

    s = new_slide(prs, "Dos componentes conservan el 58,9% de la varianza", [
        "PC1 explica 49,7% y PC2 9,2%; seis componentes superan 80% y diez superan 90%.",
        "80% y 90% son referencias descriptivas elegidas, no umbrales de selección automática.",
        "El biplot omite 41,1%. Mantenemos las 27 variables estandarizadas para SOM y Oja.",
        "El enunciado pide agrupar considerando el conjunto de variables; PCA ayuda a interpretar la representación.",
        source_pca])
    picture(s, "pca_variance.png", y=1.55, h=5.8)
    caption(s, "Kohonen y Oja reciben las 27 variables estandarizadas")

    s = new_slide(prs, "PC1 contrapone volumen y profundidad", [
        "PC1: epígrafe +0,256; impresiones -0,251; video -0,249; alcance -0,248; tendencias -0,246.",
        "PC2: showcase +0,419 y arte +0,401 frente a audiencia argentina -0,361 y perros -0,261.",
        "Las flechas muestran cinco variables interpretables, con un factor común de escala visual; no son predictores comerciales.",
        "El signo del autovector es una convención. Los puntos cercanos en dos PC pueden diferir en las restantes.",
        source_pca])
    picture(s, "pca_biplot.png", y=1.4, h=6.8)

    SECTION = "1.2"
    s = new_slide(prs, "Kohonen organiza perfiles por semejanza", [
        "Entrada: las mismas 27 variables estandarizadas. Ganadora: menor distancia euclídea.",
        "Actualizamos la ganadora y vecinas del radio binario; el dibujo destaca una ganadora y su vecindad.",
        "Grilla hexagonal, pesos iniciales con muestras; η(i)=0,9/(1+i), i cuenta presentaciones; R=1 constante.",
        "Incluimos el borde d<=R para que R=1 alcance vecinas pegadas, como los ejemplos de clase pp. 15 y 31.",
        "Los nombres y el tópico dominante solo rotulan resultados; los porcentajes de tópicos sí son variables de entrada.",
        source_som])
    picture(s, "som_neighborhood.png", x=.9, y=1.5, w=8.2, h=6.3)
    text(s, "27 variables", 10.0, 2.0, 4.9, .7, size=32)
    text(s, "w ← w + η(x − w)", 9.3, 3.6, 6.0, 1.2, size=32, color=TEAL)
    text(s, "Ganadora y vecinas", 9.8, 5.35, 5.2, .7, size=26)

    for title, filename, reminders in [
        ("Una tasa pequeña deja vecinas menos parecidas", "som_eta.png", [
            "Variamos solo η₀: 0,1, 0,5, 0,9; resto fijo en la base común.",
            "η₀=0,1 da menor distancia al ganador, pero U media mucho mayor: prototipos vecinos menos parecidos.",
            "η₀=0,9 equilibra distancia y continuidad frente a 0,5. La elección es manual y condicional a esta base."]),
        ("Un radio mayor mezcla más los perfiles", "som_radius.png", [
            "Variamos solo el radio constante: 1, 2 y 3; mantenemos tamaño, η y épocas.",
            "Al aumentar el radio suben la distancia media al ganador y la fracción de neuronas vacías.",
            "U cae, pero esa caída no prueba calidad: actualizaciones amplias acercan los prototipos entre sí."]),
        ("Más neuronas reducen la agregación", "som_size.png", [
            "Variamos solo k de la grilla k×k: 3,4,5,6; cols=null acompaña el tamaño cuadrado.",
            "Las grillas grandes reducen distancia al ganador, pero tienen más neuronas vacías y menor agregación de creadores.",
            "4×4 mantiene grupos legibles. El tamaño no tiene un criterio demostrado en la clase; no afirmamos óptimo global."]),
    ]:
        s = new_slide(prs, title, reminders + [
            "Cada valor usa 30 semillas, 42 a 71. Barras: media ± desvío muestral entre entrenamientos, ddof=1.",
            "Diseño OAT: cambia un parámetro por vez; las semillas y las demás condiciones son comunes.",
            source_som + " Fuente numérica: runs/som_sweep/comparison.csv."])
        picture(s, filename, y=1.7, h=5.8)
        caption(s, "Un parámetro por vez · 30 semillas · media ± DE")

    s = new_slide(prs, "Contrastamos duración, topología e inicialización", [
        "Cada panel varía un único parámetro sobre la misma base, no combina los factores.",
        "200 épocas cambian poco la distancia frente a 50; el descenso de η limita ajustes tardíos.",
        "Rectangular aproxima algo mejor las entradas; hexagonal tiene menor U media y menos vacías en este análisis.",
        "Muestras supera la inicialización uniforme en distancia y ocupación; preserva representantes próximos a datos.",
        "Final: 4×4 hexagonal, muestras, η₀=0,9, R=1, 200 épocas. Distancia 3,45829±0,07037; vacías 0,9±0,8847 de 16.",
        source_som + " Datos: runs/som_sweep/comparison.csv y som_design_sweep/comparison.csv."])
    picture(s, "som_design.png", y=1.7, h=5.5)
    caption(s, "4×4 · hexagonal · muestras · η₀=0,9 · R=1 · 200 épocas")

    profiles = data("som", "creator_profiles.csv").set_index("bmu")
    s = new_slide(prs, "La grilla reúne perfiles comerciales distintos", [
        "Este mapa ilustra la semilla 42; las coordenadas no identifican perfiles permanentes entre entrenamientos.",
        "BMU 5 reúne 10 creadores de mayor profundidad; BMU 13 reúne 9 de mayor volumen. Revisar medias del cuadro.",
        "BMU 1 reúne dos creadores con 54,47% de arte y PC2 media positiva; el biplot permite interpretar el contraste.",
        "El número grande es cantidad de creadores; el pequeño es índice de neurona, desde cero.",
        "Los perfiles son medias reales de los asignados; los pesos prototipo no tienen por qué coincidir con esas medias.",
        source_som + " Datos: runs/som/{neurons,creator_profiles,bmus}.csv."])
    picture(s, "som_hits.png", x=.7, y=1.5, w=8.7, h=6.5)
    for y, cell, label in [(2.0, 5, "Profundidad"), (3.8, 13, "Volumen"), (5.6, 1, "Arte")]:
        r = profiles.loc[cell]
        text(s, f"Celda {cell} · {label}", 10, y, 5.2, .6, size=28, bold=True)
        component, score = ("PC2", r.pc2) if cell == 1 else ("PC1", r.pc1)
        text(s, f"{int(r.n_creators)} creadores · {component} {score:+.2f}", 10, y + .65, 5.2, .6, size=23, color=TEAL)
    caption(s, "Ejemplo: semilla 42")

    edge = data("som", "u_edges.csv").iloc[0]
    s = new_slide(prs, "La matriz U localiza transiciones de perfil", [
        "Cada número grande es la distancia media del prototipo a sus vecinos; el pequeño es el índice de neurona.",
        f"En la semilla 42 la mayor arista une {int(edge.neuron_a)} y {int(edge.neuron_b)}; distancia {edge.distance:.4f}. Revisar top_diff en u_edges.csv.",
        "La arista resaltada es una distancia entre dos prototipos; no confundirla con la media U de una celda.",
        "U media entre 30 entrenamientos: 2,10844±0,19310. No promediamos pesos ni celdas que pueden rotarse.",
        "Distancia alta sugiere cambio de perfil; no establece una frontera estadística ni calidad comercial.",
        source_som + " Datos: runs/som/{neurons,u_edges,per_seed}.csv."])
    picture(s, "som_u.png", x=.9, y=1.5, w=9.5, h=6.4)
    text(s, "Uᵢ = media de ‖wᵢ − wⱼ‖", 10.0, 3.0, 5.0, 1.4, size=29, color=TEAL)
    text(s, "j: vecinas de i", 10.0, 4.8, 4.8, .7, size=23)
    caption(s, "Mapa de ejemplo: semilla 42")

    SECTION = "1.3"
    s = new_slide(prs, "Oja aproxima la primera componente", [
        "Una neurona: y=w·x; actualización Δw=ηy(x−yw); pesos iniciales uniformes [0,1], η=0,001, 200 épocas.",
        "Mostramos seis coeficientes interpretables; pesos completos y dispersiones en weights_stats.csv.",
        "Coseno absoluto con PC1: media 0,999891, DE 0,000039 en 30 entrenamientos.",
        "PC1 solo permite comparar y orientar el informe. No interviene en actualización ni parada.",
        "Norma cerca de uno no prueba dirección convergida. Con η constante persiste variación del vector.",
        "Fuente: clase Oja y Sanger pp. 12-16; runs/oja/{repeats,weights_stats,history_stats}.csv."])
    picture(s, "oja_weights.png", x=.7, y=1.6, w=10.5, h=6.4)
    text(s, "0,999891", 11.5, 2.55, 4, .9, size=43, color=TEAL, bold=True)
    text(s, "± 0,000039", 11.5, 3.6, 4, .7, size=26)
    text(s, "Coseno con PC1\n30 semillas", 11.5, 4.8, 4, 1.4, size=23, color=MUTED)

    s = new_slide(prs, "El índice ordena un eje de variación", [
        "Cada punto es un creador; eje Y muestra media y DE del índice entre semillas.",
        "Normalizamos y alineamos el signo de los vectores aprendidos para informar; nunca renormalizamos durante el entrenamiento.",
        "Los cinco primeros aparecen en top5 en los 30 entrenamientos; las cinco cuentas corresponden al extremo positivo de PC1.",
        "Índice alto no mide ventas, adopciones ni adecuación universal. Invertir el signo invierte el ranking.",
        "La decisión necesita el objetivo de campaña y las variables observadas, no el máximo automático.",
        "Fuente: clase PCA p.46; Oja pp.12-16; runs/oja/index_stats.csv y ranking.csv."])
    picture(s, "oja_scores.png", x=.8, y=1.65, w=10.8, h=6.1)
    text(s, "Top 5 estable", 11.7, 2.65, 3.7, .8, size=29, bold=True)
    text(s, "30/30", 11.7, 3.75, 3.7, .9, size=44, color=TEAL)
    text(s, "entrenamientos", 11.7, 4.9, 3.7, .7, size=23)

    SECTION = "1.4"
    s = new_slide(prs, "Para fútbol priorizamos profundidad", [
        "Supuesto: explicar ropa de entrenamiento; filtro de contenido de fútbol >=40%, orden por guardados por 1000.",
        "Comparamos solo candidatos de fútbol de cada celda. No usar medias de toda la celda como si fueran ese subconjunto.",
        "Celda 5 agrupa los tres elegidos de fútbol; celda 13 concentra más volumen y menos profundidad educativa.",
        "Una campaña de reconocimiento podría priorizar alcance o descubrimiento; el objetivo cambia la decisión.",
        "No hay ventas, conversiones ni adopciones en el catálogo: estos proxies no prueban rentabilidad.",
        "Fuente: enunciado 1.4; configs/campaigns.json; runs/campaigns/strategy_profiles.csv."])
    picture(s, "football_profiles.png", y=1.7, h=5.7)
    caption(s, "Candidatos con ≥40% de fútbol · ejemplo de la semilla 42")

    s = new_slide(prs, "Cada campaña exige su propia selección", [
        "Supuestos comunes: una publicación por creador, tres creadores y contenido del tópico >=40%.",
        "Perros ordena por visitas al perfil; fútbol por guardados; belleza y arte por proporción de tutoriales.",
        "Comparación: solo seguidores globales y seguidores dentro del tópico. El control temático separa el efecto del filtro.",
        "La selección propuesta pierde alcance y a veces especialización; en estos datos obtiene menor precio y más profundidad según el criterio.",
        "La suma de alcances puede repetir audiencia. Los supuestos comerciales se revisan con resultados reales.",
        "Fuente: runs/campaigns/{picks,baselines,comparison,evidence}.csv; configs/campaigns.json."])
    for y, client in zip([2.0, 3.35, 4.7, 6.05], ["dogs", "football", "makeup", "art"]):
        text(s, TOPICS[client], .9, y, 2.4, .65, size=29, color=TEAL, bold=True)
        handles = campaigns[f"{client}_picks"]
        for x, handle in zip([3.6, 7.6, 11.6], handles):
            text(s, handle, x, y + .08, 4.0, .75, size=19)

    s = new_slide(prs, "Los cambios de datos obligan a revisar el análisis", [
        "Campaña nueva con mismo catálogo: definir objetivo y revisar evidencia; puede cambiar la elección comercial sin cambiar los modelos.",
        "Datos mensuales o más creadores: revisar distribuciones, escalas, entrenamientos, grilla e interpretaciones.",
        "Variables nuevas cambian distancias y componentes: repetir representación, estandarización y modelos.",
        "Puede proyectarse un creador nuevo con la transformación existente si sigue siendo representativa; el repositorio regenera todo al actualizar el CSV.",
        "Guardar versión de datos, fecha, semillas y parámetros; celdas y signos no son etiquetas permanentes.",
        "Fuente: enunciado 1.4; modelo de Kohonen y proyección PCA; configs y runs."])
    for y, left, right in [(2.4, "Campaña nueva", "Revisar objetivo y criterios"),
                           (4.1, "Datos actualizados", "Revisar escalas y entrenamientos"),
                           (5.8, "Variables nuevas", "Repetir representación y modelos")]:
        text(s, left, .95, y, 5.2, .8, size=30, bold=True)
        text(s, "→", 6.45, y, 1, .8, size=34, color=MUTED)
        text(s, right, 8.0, y, 7, 1.0, size=28, color=TEAL)

    SECTION = "2.1"
    s = new_slide(prs, "Primero almacenamos cuatro letras", [
        "Cada letra 5×5 se recorre por filas: 25 estados bipolares. Negro=+1, blanco=-1.",
        "G, R, T, V forman cuatro filas de P: matriz 4×25. Cada posición corresponde a una neurona.",
        "Hebb: W=PᵀP/25; W es 25×25 y su diagonal se fija en cero. Los cuatro patrones se almacenan juntos.",
        "Un patrón está almacenado por construcción de W; estar estable y recuperar desde ruido se comprueba después.",
        "Ejemplo introductorio antes del caso QR; fuente de patrones: datos/letras_5x5.json.",
        source_hop + " Figura original de esta corrida: runs/hopfield_letters/storage.png."])
    picture(s, "hopfield_letters__storage.png", x=.65, y=1.6, w=14.8, h=6.35)

    s = new_slide(prs, "La semejanza entre recuerdos cambia su estabilidad", [
        "Productos internos normalizados: (ξᵘ·ξᵛ)/(5×5), antes de agregar ruido. Norma de cada letra: 5.",
        "N=25 es impar: el producto interno de dos vectores bipolares no puede ser cero; mínimo absoluto posible: 1.",
        "Elegimos G, R, T, V minimizando máximo |producto|, luego promedio; producto bruto absoluto 1, normalizado 0,04.",
        "H, M, N, W tienen máximo absoluto 21/25=0,84 y solo una letra estable, almacenando la misma cantidad.",
        "La diagonal normalizada 1 se excluye del resumen entre patrones distintos; la figura conserva el signo.",
        source_hop + " Datos: dot_products.csv, orthogonality.csv, stored.csv."])
    picture(s, "letters_dots.png", y=1.55, h=6.1)
    caption(s, "G,R,T,V: 4/4 estables · H,M,N,W: 1/4")

    s = new_slide(prs, "La recuperación cae al aumentar el ruido", [
        "Cada semilla contiene 50 consultas por letra y nivel. Se invierten posiciones distintas dentro de cada consulta.",
        "G, R, T, V asincrónica: 3 bits, 99,833% ± 0,240 pp; 5 bits, 93,417% ± 2,060 pp; 9 bits, 31,917% ± 2,910 pp.",
        "Media y DE muestral entre 30 tasas por semilla. Las consultas de una corrida se agregan antes de calcular dispersión.",
        "192.000 consultas de la grilla estándar; 24.000 adicionales para nivel 0,4: 216.000 consultas totales.",
        "Actualización sincrónica puede formar ciclos; asincrónica tiene energía no creciente. No prometer recuperación para toda consulta.",
        source_hop + " Datos: recall_by_seed.csv y recall_repetitions.csv."])
    picture(s, "letters_noise.png", y=1.65, h=6.0)
    caption(s, "30 semillas · media ± DE · 50 consultas por letra y nivel")

    s = new_slide(prs, "La energía puede bajar hacia un espurio", [
        "Ejemplo de la semilla 42: consulta de G con 10/25 bits invertidos.",
        "El estado final es un punto fijo ajeno a las cuatro letras y sus inversas.",
        "Energía final -7,68; distancia mínima a una letra: 6 bits. Los recuerdos guardados tienen energía -10,56.",
        "La curva registra visitas individuales a neuronas, incluida la comprobación final. Un espurio puede tener energía estable.",
        "Distinguir convergencia a un punto fijo de recuperación del patrón objetivo.",
        source_hop + " Datos: summary.json y energy_trace.csv, ejemplo spurious."])
    picture(s, "letters_spurious.png", y=2.0, h=5.6)
    caption(s, "Consulta de G con 10 bits invertidos · semilla 42")

    s = new_slide(prs, "La capacidad depende también de los patrones", [
        "Se prueban 200 subconjuntos aleatorios del alfabeto por cantidad, en 30 semillas.",
        "La curva muestra proporción de subconjuntos enteramente estables, con media ± DE entre semillas.",
        "Para 4 letras al azar: 48,35% ± 3,57 pp de conjuntos enteramente estables; la fracción de letras estables es 75,575%.",
        "0,15N=3,75 es una aproximación para patrones poco correlacionados, sin garantía para cada conjunto de 25 neuronas.",
        "Las cuatro letras elegidas son estables pese a superar 3,75; compatible con una capacidad aproximada.",
        source_hop + " Datos: capacity_by_seed.csv y capacity_repetitions.csv."])
    picture(s, "letters_capacity.png", y=1.6, h=6.0)
    caption(s, "200 subconjuntos por cantidad · 30 semillas · media ± DE")

    s = new_slide(prs, "Usamos los cinco QR originales de 37×37", [
        "Las cinco imágenes coinciden byte a byte con el ZIP suministrado; inputs.csv registra SHA256, dimensiones y norma.",
        "Cada imagen 37×37 se recorre por filas. Negro=+1, blanco=-1; se conservan todos sus módulos.",
        "Cada vector tiene 1369 estados y norma √1369=37. P es 5×1369 y W es 1369×1369, con diagonal cero.",
        "La comparación usa coseno normalizado; la red utiliza estados bipolares ±1.",
        "El render de esta diapositiva solo amplía píxeles por vecino más próximo; no altera los archivos originales.",
        "Fuente: enunciado 2.1; datos/hopfield/*.png; runs/hopfield_qr/inputs.csv."])
    picture(s, "qr_originals.png", y=1.75, h=4.7)
    text(s, "37×37 → 1369 estados ±1", .95, 6.65, 14.2, .7, size=36, color=TEAL)
    text(s, "P: 5×1369     W: 1369×1369", 1, 7.65, 14, .6, size=25)

    pairs = data("hopfield_qr", "pairwise_similarity.csv")
    cosine = float(pairs.normalized_dot.mean())
    s = new_slide(prs, "Los QR comparten una similitud alta", [
        f"Coseno medio entre los 10 pares distintos: {cosine:.6f}; rango {pairs.normalized_dot.min():.6f} a {pairs.normalized_dot.max():.6f}.",
        "Para ±1 y N=1369, cada norma es 37: el denominador es 37×37=1369.",
        f"1−coseno medio={1-cosine:.6f}; 0,3299 es disimilitud complementaria, una cantidad diferente del producto normalizado.",
        "La fracción de bits distintos es (1−coseno)/2: aproximadamente 16,49% en promedio.",
        "955/1369 módulos son iguales en los cinco QR simultáneamente; esa proporción, 69,76%, es otra cantidad.",
        "Fuente: runs/hopfield_qr/{pairwise_similarity,overlaps,inputs}.csv; clase Hopfield p.33."])
    picture(s, "qr_similarity.png", x=.8, y=1.5, w=8.5, h=6.6)
    text(s, f"{cosine:.4f}".replace(".", ","), 10, 2.45, 5, 1.1, size=62, color=TEAL, bold=True)
    text(s, "Similitud media", 10, 3.9, 5, .7, size=27)
    text(s, "sᵘᵛ = (ξᵘ·ξᵛ) / (37·37)", 9.4, 5.45, 6, 1.2, size=27)

    s = new_slide(prs, "Cinco QR juntos pierden el recuerdo consultado", [
        "Con los cinco juntos, ningún original es punto fijo: 71, 91, 225, 225 y 103 módulos inestables.",
        "33.000 consultas conjuntas: 5 QR × 11 niveles × 20 consultas × 30 semillas. Recuperación exacta: 0% en todos los niveles.",
        "Ejemplo de la semilla 42: original principal sin ruido y estado final espurio. La curva de barridos usa 30 semillas.",
        "Máximo 2 barridos que cambian el estado; la comprobación final no cuenta como otro cambio.",
        "Energía asincrónica no creciente no garantiza recuerdo correcto. El fracaso aparece ya en estabilidad, antes del ruido.",
        "Fuente: clase Hopfield pp.34-38; runs/hopfield_qr/{stored,recall_repetitions,queries}.csv."])
    picture(s, "qr_failure.png", y=1.7, h=5.5)
    caption(s, "0% de recuperación exacta · 33.000 consultas conjuntas")

    s = new_slide(prs, "Al reducir la carga, estos QR sí se recuperan", [
        "Controles de estabilidad exhaustivos: 5 redes con un QR y 10 parejas; todos sus patrones son estables.",
        "Los 10 tríos, 5 cuartetos y la red de cinco tienen cero patrones originales estables.",
        "Un QR por red: 33.000 consultas adicionales; recuperación 100% hasta 684/1369 bits invertidos (49,963%), con DE 0 entre semillas.",
        "El fallo no es universal para QR ni atribuible solo a su tamaño. Estos patrones correlacionados interfieren cuando se almacenan juntos.",
        "El control cambia la matriz de pesos. No comparar energías absolutas entre redes como si fueran una misma función.",
        "Fuente: runs/hopfield_qr/{storage_control,single_pattern_control,recall_repetitions}.csv."])
    picture(s, "qr_controls.png", y=1.75, h=5.9)
    caption(s, "Mismos originales, distinta cantidad almacenada por red")

    s = new_slide(prs, "Recuperar el patrón y leer la URL son pruebas distintas", [
        "El lector opcional decodifica 5/5 originales. Se compara la URL con la del patrón objetivo.",
        "En esta prueba el lector directo logra 100% sin ruido y 0% desde 5%; la salida conjunta de Hopfield obtiene 0%.",
        "Los porcentajes son medias sobre 30 semillas; las barras DE 0 reflejan estos resultados observados, no una garantía.",
        "Las entradas ya son matrices alineadas 37×37. Fotografías reales agregarían localización, perspectiva y binarización.",
        "Una carta o QR nuevos requieren actualizar memorias y volver a comprobar estabilidad e interferencia.",
        "Fuente: enunciado 2.1; runs/hopfield_qr/recall_repetitions.csv y summary.json; lector zxing-cpp."])
    picture(s, "qr_reader.png", y=1.6, h=6.2)
    caption(s, "Matrices alineadas de 37×37; fotografías reales quedan fuera de esta prueba")

    prs.save(OUT)
    print(f"{OUT}: {len(prs.slides)} diapositivas, {sum(bool(s.notes_slide.notes_text_frame.text.strip()) for s in prs.slides)} con notas")


if __name__ == "__main__":
    build()
